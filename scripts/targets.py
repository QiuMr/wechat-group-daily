#!/usr/bin/env python3
"""Build the monthly publisher target dashboard from daily facts.json files."""

from __future__ import annotations

import argparse
import calendar
import json
import re
from collections import Counter
from datetime import date as date_cls, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from analytics import contact_key
from common import SKILL, data_home, date_dir, now_text, read_json, write_json
from report import style_sheet

STATUS_LABELS = {
    "done": "已达标",
    "ontrack": "跟上",
    "behind": "落后",
    "critical": "严重落后",
    "unset": "未设目标",
}


def _norm_name(value: str) -> str:
    return re.sub(r"\s+", "", str(value or "")).lower()


def _num(value: Any, default: float = 0.0) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _pct(value: float | None) -> float | None:
    return round(value, 1) if value is not None else None


def _status(completion: float | None, pace_diff: float | None, target: float) -> str:
    if not target:
        return "unset"
    if completion is None:
        return "behind"
    if completion >= 100:
        return "done"
    if pace_diff is not None and pace_diff >= 0:
        return "ontrack"
    if pace_diff is not None and pace_diff >= -10:
        return "behind"
    return "critical"


class TargetBook:
    def __init__(self, raw: dict[str, Any]) -> None:
        self.raw = raw
        self.days = max(1, min(31, int(raw.get("days") or 30)))
        self.month = str(raw.get("month") or "").strip()
        self.publishers: list[dict[str, Any]] = []
        self.aliases: dict[str, str] = {}
        self.warnings: list[str] = []
        seen: set[str] = set()
        for item in raw.get("publishers") or []:
            if not isinstance(item, dict):
                continue
            name = str(item.get("name") or "").strip()
            if not name or name in seen:
                continue
            seen.add(name)
            target = max(0.0, _num(item.get("monthly_target")))
            publisher = {
                "name": name,
                "monthly_target": target,
                "aliases": [name] + [str(alias).strip() for alias in item.get("aliases") or [] if str(alias).strip()],
            }
            self.publishers.append(publisher)
            for alias in publisher["aliases"]:
                key = _norm_name(alias)
                if key and key not in self.aliases:
                    self.aliases[key] = name
        if not self.publishers:
            self.warnings.append("targets.json 里没有配置任何发布员目标")
        self.names = [item["name"] for item in self.publishers]
        self.by_name = {item["name"]: item for item in self.publishers}

    def resolve(self, raw_name: str) -> str:
        name = str(raw_name or "").strip()
        if not name:
            return "未知发布员"
        return self.aliases.get(_norm_name(name), name)


def load_targets(path: Path) -> TargetBook:
    raw = read_json(path)
    if not isinstance(raw, dict):
        raise ValueError(f"读不到目标配置：{path}")
    return TargetBook(raw)


def _month_start(month: str) -> date_cls:
    try:
        return datetime.strptime(f"{month}-01", "%Y-%m-%d").date()
    except ValueError as exc:
        raise ValueError(f"月份格式应为 YYYY-MM：{month}") from exc


def _facts_for(date_value: date_cls, home: Path) -> tuple[dict[str, Any] | None, str]:
    path = date_dir(date_value, home) / "report" / "facts.json"
    data = read_json(path)
    if not isinstance(data, dict):
        return None, "missing"
    customers = data.get("customers")
    if not isinstance(customers, list):
        return data, "legacy"
    return data, "ok"


def build_target_data(
    month: str,
    report_date: date_cls,
    book: TargetBook,
    home: Path,
    current_facts: Path | None = None,
) -> dict[str, Any]:
    month_start = _month_start(month)
    month_days = calendar.monthrange(month_start.year, month_start.month)[1]
    month_end = month_start.replace(day=month_days)
    if not (month_start <= report_date <= month_end):
        raise ValueError(f"统计日 {report_date} 不在月份 {month} 内")
    target_days = book.days
    day_index = min(report_date.day, target_days)
    time_progress = round(min(day_index / target_days * 100, 100.0), 1)
    remaining_days = max(target_days - day_index, 0)
    if report_date >= month_end:
        remaining_days = 0

    configured = set(book.names)
    seen_contacts: set[str] = set()
    cumulative: Counter[str] = Counter()
    daily_rows: list[dict[str, Any]] = []
    warnings = list(book.warnings)
    missing_days: list[int] = []
    legacy_days: list[int] = []

    for order in range(1, report_date.day + 1):
        current = month_start.replace(day=order)
        if current == report_date and current_facts and current_facts.exists():
            data = read_json(current_facts)
            state = "ok" if isinstance(data, dict) and isinstance(data.get("customers"), list) else "legacy"
        else:
            data, state = _facts_for(current, home)
        row = {
            "date": current.isoformat(),
            "day": order,
            "today": None,
            "today_configured": None,
            "today_unconfigured": None,
            "by_publisher": {},
            "state": state,
        }
        if state == "missing":
            missing_days.append(order)
        elif state == "legacy":
            legacy_days.append(order)
        else:
            counter: Counter[str] = Counter()
            for customer in data.get("customers") or []:
                if not isinstance(customer, dict):
                    continue
                contact = customer.get("contact") if isinstance(customer.get("contact"), dict) else {}
                key = contact_key(contact)
                if not key or key in {"phone:", "wechat:"} or key in seen_contacts:
                    continue
                seen_contacts.add(key)
                publisher = book.resolve(str(customer.get("publisher") or ""))
                counter[publisher] += 1
                cumulative[publisher] += 1
            row["by_publisher"] = dict(counter)
            row["today"] = sum(counter.values())
            row["today_configured"] = sum(count for name, count in counter.items() if name in configured)
            row["today_unconfigured"] = sum(count for name, count in counter.items() if name not in configured)
        daily_rows.append(row)

    if missing_days:
        warnings.append(
            "缺日报：" + "、".join(f"{month_start.month}-{day:02d}" for day in missing_days[:12])
            + (" 等" if len(missing_days) > 12 else "")
            + "；月累计按已知数据计算，可能偏低"
        )
    if legacy_days:
        warnings.append(
            "旧版数据没有客户明细：" + "、".join(f"{month_start.month}-{day:02d}" for day in legacy_days[:12])
            + (" 等" if len(legacy_days) > 12 else "")
            + "；这些日期未计入目标"
        )

    team_target = sum(item["monthly_target"] for item in book.publishers)
    last_row = daily_rows[-1] if daily_rows else None
    last_has_data = bool(last_row and last_row["state"] == "ok")
    team_completed = sum(cumulative[name] for name in configured)
    team_today = last_row["today_configured"] if last_has_data else None
    team_today_all = last_row["today"] if last_has_data else None
    unconfigured_completed = sum(count for name, count in cumulative.items() if name not in configured)
    unconfigured_today = last_row["today_unconfigured"] if last_has_data else None

    team_completion = round(team_completed * 100 / team_target, 1) if team_target else 0.0
    team_pace = round(team_completion - time_progress, 1)
    team_daily_target = round(team_target / target_days, 1) if target_days else 0.0
    team_today_completion = (
        round(team_today * 100 / team_daily_target, 1)
        if team_today is not None and team_daily_target
        else None
    )
    team_gap = max(round(team_target - team_completed, 1), 0.0)
    elapsed = max(report_date.day, 1)
    if team_completed and elapsed:
        forecast = round(team_completed / elapsed * target_days)
    else:
        forecast = 0
    if report_date >= month_end or report_date.day >= target_days:
        forecast = team_completed
    team_status = _status(team_completion, team_pace, team_target)

    champion: list[dict[str, Any]] = []
    for item in book.publishers:
        name = item["name"]
        target = item["monthly_target"]
        completed = cumulative[name]
        daily_target = round(target / target_days, 1) if target_days else 0.0
        today = last_row["by_publisher"].get(name, 0) if last_has_data else None
        completion = round(completed * 100 / target, 1) if target else None
        pace = round(completion - time_progress, 1) if completion is not None else None
        gap = max(round(target - completed, 1), 0.0)
        remaining_daily = round(gap / remaining_days, 1) if remaining_days and gap > 0 else None
        status = _status(completion, pace, target)
        champion.append(
            {
                "name": name,
                "monthly_target": round(target, 1),
                "daily_target": daily_target,
                "today": today,
                "today_completion": round(today * 100 / daily_target, 1)
                if today is not None and daily_target
                else None,
                "completed": completed,
                "completion": completion,
                "gap": gap,
                "remaining_daily": remaining_daily,
                "pace_diff": pace,
                "status": status,
                "status_label": STATUS_LABELS[status],
                "bar": min(completion or 0, 100),
                "configured": True,
            }
        )
    champion.sort(key=lambda row: (-(row["completion"] or 0), -row["completed"], row["name"]))
    for index, row in enumerate(champion, start=1):
        row["rank"] = index

    unconfigured = []
    for name, completed in sorted(cumulative.items(), key=lambda item: (-item[1], item[0])):
        if name in configured:
            continue
        unconfigured.append(
            {
                "name": name,
                "today": last_row["by_publisher"].get(name, 0) if last_has_data else None,
                "completed": completed,
            }
        )
    if unconfigured:
        warnings.append("未设目标的发布员：" + "、".join(item["name"] for item in unconfigured))

    if target_days > month_days:
        warnings.append(
            f"本月只有 {month_days} 天，目标却按 {target_days} 天配置；"
            f"月底时间进度最多 {round(month_days * 100 / target_days, 1)}%，"
            f"建议把 config/targets.json 的 days 改为 {month_days}"
        )
    chart_days = min(max(target_days, report_date.day), month_days)
    cumulative_by_day: dict[int, int] = {}
    running = 0
    for row in daily_rows:
        if row["state"] == "ok":
            running += row["today_configured"] or 0
        cumulative_by_day[row["day"]] = running
    last_known = 0
    daily_series = []
    for day in range(1, chart_days + 1):
        row = daily_rows[day - 1] if day <= len(daily_rows) else None
        if row and row["state"] == "ok":
            last_known = cumulative_by_day.get(day, last_known)
            today = row["today_configured"]
            today_all = row["today"]
            missing = False
            legacy = False
            future = False
        elif row and row["state"] in {"missing", "legacy"}:
            today = None
            today_all = None
            missing = row["state"] == "missing"
            legacy = row["state"] == "legacy"
            future = False
        else:
            today = None
            today_all = None
            missing = False
            legacy = False
            future = True
        daily_series.append(
            {
                "day": day,
                "date": month_start.replace(day=day).isoformat() if day <= month_days else "",
                "today": today,
                "today_all": today_all,
                "cumulative": cumulative_by_day.get(day, last_known) if not future else None,
                "ideal": round(team_target * min(day, target_days) / target_days, 1),
                "missing": missing,
                "legacy": legacy,
                "future": future,
            }
        )

    if team_target and team_completed > team_target:
        warnings.append("团队月累计已超过月目标，完成率按实际值展示")
    if not last_has_data:
        warnings.append("统计日没有可用的客户明细，今日数据和团队今日完成率为空")

    month_label = f"{month_start.year} 年 {month_start.month} 月"
    date_label = f"{month_start.month} 月 {report_date.day} 日"
    notes = [
        {"label": "统计范围", "value": f"{month_start.isoformat()}–{report_date.isoformat()}"},
        {"label": "客户口径", "value": "有联系方式；按手机号/微信号归一化去重；同月第一次出现才计数。"},
        {"label": "归属规则", "value": "同一客户同月归第一个发布员；后续重复发布不再重复计数。"},
        {"label": "日目标", "value": f"月目标 ÷ {target_days}；当前日目标按团队/个人月目标换算。"},
        {"label": "时间进度", "value": f"第 {day_index}/{target_days} 天 = {time_progress}%；按 30 天口径封顶 100%。"},
        {"label": "状态规则", "value": "完成率 ≥ 时间进度为跟上；低 10 个点以内为落后；低超过 10 个点为严重落后。"},
        {"label": "未设目标", "value": "只展示实际产出，不计入团队目标、完成率和排名。"},
        {"label": "数据来源", "value": "每日 report/facts.json；缺日报的日期标灰，累计按已知数据计算。"},
    ]

    return {
        "meta": {
            "title": f"{month_start.month} 月 · 发布员目标进度",
            "month": f"{month_start.year:04d}-{month_start.month:02d}",
            "month_label": month_label,
            "date": report_date.isoformat(),
            "date_label": date_label,
            "day_index": day_index,
            "target_days": target_days,
            "time_progress": time_progress,
            "remaining_days": remaining_days,
            "generated_at": now_text().replace("T", " ")[:16],
            "export_name": f"目标进度_{month_start.year:04d}-{month_start.month:02d}.xlsx",
            "brand_sub": f"{month_label} · {len(book.publishers)} 位发布员",
            "eyebrow": f"{month_label} · 第 {day_index}/{target_days} 天",
            "scope_note": f"统计范围 {month_start.month}-{month_start.day:02d}–{report_date.month}-{report_date.day:02d}",
        },
        "team": {
            "target": round(team_target, 1),
            "completed": team_completed,
            "completion": team_completion,
            "gap": team_gap,
            "today": team_today,
            "today_all": team_today_all,
            "daily_target": team_daily_target,
            "today_completion": team_today_completion,
            "pace_diff": team_pace,
            "forecast": forecast,
            "status": team_status,
            "status_label": STATUS_LABELS[team_status],
            "configured": len(book.publishers),
            "unconfigured": len(unconfigured),
            "unconfigured_completed": unconfigured_completed,
            "unconfigured_today": unconfigured_today,
            "time_progress": time_progress,
        },
        "publishers": champion,
        "unconfigured": unconfigured,
        "daily": daily_series,
        "warnings": list(dict.fromkeys(warnings)),
        "notes": notes,
    }


def _fmt_num(value: Any, digits: int = 0) -> str:
    if value is None:
        return "—"
    if digits:
        return f"{float(value):,.{digits}f}"
    return f"{round(float(value)):,}"


def build_xlsx(data: dict[str, Any], path: Path) -> None:
    meta = data["meta"]
    team = data["team"]
    workbook = Workbook()
    workbook.properties.title = meta["title"]
    workbook.properties.creator = "wechat-group-daily"

    team_sheet = workbook.active
    team_sheet.title = "团队"
    rows = [
        ["月份", meta["month_label"]],
        ["统计日", meta["date"]],
        ["时间进度", f"第 {meta['day_index']}/{meta['target_days']} 天 · {meta['time_progress']}%"],
        ["团队月目标", _fmt_num(team["target"])],
        ["团队月累计", _fmt_num(team["completed"])],
        ["月完成率", f"{team['completion']}%"],
        ["进度差", f"{team['pace_diff']:+.1f} 个点"],
        ["距月目标", _fmt_num(team["gap"])],
        ["今日新增", _fmt_num(team["today"])],
        ["日均目标", _fmt_num(team["daily_target"], 1)],
        ["今日完成率", f"{team['today_completion']}%" if team["today_completion"] is not None else "—"],
        ["月底预计", _fmt_num(team["forecast"])],
        ["未设目标累计", _fmt_num(team["unconfigured_completed"])],
        ["状态", team["status_label"]],
    ]
    team_sheet.append(["指标", "数值"])
    for row in rows:
        team_sheet.append(row)
    for cell in team_sheet[1]:
        cell.fill = PatternFill("solid", fgColor="0E7C66")
        cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_index in range(2, team_sheet.max_row + 1):
        for column_index in range(1, 3):
            cell = team_sheet.cell(row=row_index, column=column_index)
            cell.font = Font(name="微软雅黑", size=10, color="1F2826")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        team_sheet.row_dimensions[row_index].height = 24
    team_sheet.column_dimensions["A"].width = 18
    team_sheet.column_dimensions["B"].width = 42
    team_sheet.freeze_panes = "A2"
    team_sheet.sheet_view.showGridLines = False
    team_sheet.sheet_properties.tabColor = "0E7C66"

    publisher_sheet = workbook.create_sheet("发布员目标")
    publisher_rows = []
    for item in data["publishers"]:
        publisher_rows.append(
            [
                item["rank"],
                item["name"],
                _fmt_num(item["monthly_target"]),
                _fmt_num(item["daily_target"], 1),
                _fmt_num(item["today"]),
                f"{item['today_completion']}%" if item["today_completion"] is not None else "—",
                _fmt_num(item["completed"]),
                f"{item['completion']}%" if item["completion"] is not None else "—",
                f"{meta['time_progress']}%",
                f"{item['pace_diff']:+.1f}" if item["pace_diff"] is not None else "—",
                _fmt_num(item["gap"]),
                _fmt_num(item["remaining_daily"], 1) if item["remaining_daily"] is not None else "—",
                item["status_label"],
            ]
        )
    for item in data["unconfigured"]:
        publisher_rows.append(
            [
                "—",
                item["name"],
                "—",
                "—",
                _fmt_num(item["today"]),
                "—",
                _fmt_num(item["completed"]),
                "—",
                f"{meta['time_progress']}%",
                "—",
                "—",
                "—",
                "未设目标",
            ]
        )
    style_sheet(
        publisher_sheet,
        ["排名", "发布员", "月目标", "日均目标", "今日新增", "今日完成率", "月累计", "月完成率", "时间进度", "进度差", "距月目标", "剩余日均", "状态"],
        [7, 20, 10, 10, 10, 12, 10, 11, 10, 9, 10, 10, 11],
        publisher_rows,
    )
    publisher_sheet.sheet_properties.tabColor = "6D5BD0"

    daily_sheet = workbook.create_sheet("每日明细")
    daily_rows = []
    for item in data["daily"]:
        daily_rows.append(
            [
                item["date"] or f"第 {item['day']} 天",
                item["day"],
                "缺日报" if item["missing"] else ("未到" if item["future"] else "有数据"),
                _fmt_num(item["today"]),
                _fmt_num(item["today_all"]),
                _fmt_num(item["cumulative"]),
                _fmt_num(item["ideal"]),
            ]
        )
    style_sheet(
        daily_sheet,
        ["日期", "第几天", "状态", "团队今日新增", "全部今日新增", "团队累计", "目标节奏"],
        [14, 8, 10, 12, 12, 12, 12],
        daily_rows,
    )
    daily_sheet.sheet_properties.tabColor = "2F6F9F"

    notes_sheet = workbook.create_sheet("口径与提醒")
    notes_sheet.append(["类型", "内容"])
    for note in data["notes"]:
        notes_sheet.append(["口径", f"{note['label']}：{note['value']}"])
    for warning in data["warnings"]:
        notes_sheet.append(["提醒", warning])
    for cell in notes_sheet[1]:
        cell.fill = PatternFill("solid", fgColor="B45309")
        cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_index in range(2, notes_sheet.max_row + 1):
        for column_index in range(1, 3):
            cell = notes_sheet.cell(row=row_index, column=column_index)
            cell.font = Font(name="微软雅黑", size=10, color="1F2826")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        notes_sheet.row_dimensions[row_index].height = 30
    notes_sheet.column_dimensions["A"].width = 10
    notes_sheet.column_dimensions["B"].width = 100
    notes_sheet.freeze_panes = "A2"
    notes_sheet.sheet_view.showGridLines = False
    notes_sheet.sheet_properties.tabColor = "6B746F"

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def build_html(data: dict[str, Any], template_path: Path, path: Path) -> None:
    template = template_path.read_text(encoding="utf-8")
    payload = json.dumps(data, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    replacements = {
        "__DATA__": payload,
        "__TITLE__": str(data["meta"]["title"]),
        "__EXPORT_NAME__": str(data["meta"]["export_name"]),
    }
    for key, value in replacements.items():
        template = template.replace(key, value)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(template, encoding="utf-8")


def build_for_date(
    report_date: date_cls,
    home: Path,
    *,
    config_path: Path | None = None,
    out_dir: Path | None = None,
    current_facts: Path | None = None,
) -> dict[str, Any]:
    config = config_path or (SKILL / "config" / "targets.json")
    if not config.exists():
        declined = (config.parent / "targets.off").exists()
        return {
            "state": "skipped",
            "reason": "no_targets_config",
            "config": str(config),
            "message": "还没有配置发布员月目标",
            "onboarding": not declined,
            "next": "日报已完成；用户没关掉引导时，可以问一次要不要设置月目标；设置后重算看板",
        }
    book = load_targets(config)
    if book.month and book.month != report_date.strftime("%Y-%m"):
        return {
            "state": "skipped",
            "reason": "month_mismatch",
            "config_month": book.month,
            "report_month": report_date.strftime("%Y-%m"),
            "message": "目标配置的月份和本次统计月份不一致",
            "onboarding": False,
            "next": "确认本月目标后更新 config/targets.json 的 month，或删除 month 字段",
        }
    data = build_target_data(
        report_date.strftime("%Y-%m"),
        report_date,
        book,
        home,
        current_facts=current_facts,
    )
    folder = out_dir or (date_dir(report_date, home) / "report")
    folder.mkdir(parents=True, exist_ok=True)
    html_path = folder / "目标进度.html"
    xlsx_path: Path | None = folder / data["meta"]["export_name"]
    facts_path = folder / "target_facts.json"
    output_warnings: list[str] = []
    try:
        build_xlsx(data, xlsx_path)
    except OSError as exc:
        fallback = folder / f"{xlsx_path.stem}_{datetime.now():%H%M%S}.xlsx"
        try:
            build_xlsx(data, fallback)
            xlsx_path = fallback
            output_warnings.append(f"目标 Excel 原文件被占用，已另存为 {fallback.name}")
        except OSError:
            xlsx_path = None
            output_warnings.append(f"目标 Excel 生成失败（可能被占用）：{exc}；HTML 已正常生成")
    if xlsx_path is not None:
        data["meta"]["export_name"] = xlsx_path.name
    build_html(data, SKILL / "assets" / "target_template.html", html_path)
    write_json(facts_path, data)
    return {
        "state": "ready",
        "month": data["meta"]["month"],
        "date": data["meta"]["date"],
        "team": data["team"],
        "html": str(html_path.resolve()),
        "xlsx": str(xlsx_path.resolve()) if xlsx_path else "",
        "facts": str(facts_path.resolve()),
        "warnings": list(dict.fromkeys([*data["warnings"], *output_warnings])),
        "top": data["publishers"][:3],
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="生成发布员目标进度看板")
    parser.add_argument("--month", default="", help="YYYY-MM；默认取 --date 所在月")
    parser.add_argument("--date", default="yesterday", help="统计日")
    parser.add_argument("--config", default=str(SKILL / "config" / "targets.json"))
    parser.add_argument("--home", default="")
    parser.add_argument("--out", default="", help="输出目录；默认 <日期>/report")
    args = parser.parse_args()

    from common import resolve_date

    value = resolve_date(args.date)
    month = args.month or value.strftime("%Y-%m")
    if month != value.strftime("%Y-%m"):
        print(json.dumps({"state": "failed", "error": "月份和统计日不一致"}, ensure_ascii=False))
        return 2
    home = Path(args.home).expanduser().resolve() if args.home else data_home()
    out_dir = Path(args.out).expanduser().resolve() if args.out else None
    try:
        result = build_for_date(value, home, config_path=Path(args.config), out_dir=out_dir)
    except (ValueError, OSError) as exc:
        print(json.dumps({"state": "failed", "error": str(exc)}, ensure_ascii=False))
        return 1
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result.get("state") == "ready" else 1


if __name__ == "__main__":
    raise SystemExit(main())
