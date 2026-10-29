#!/usr/bin/env python3
"""Build the daily group report from a probe output folder."""

from __future__ import annotations

import argparse
import html
import json
import re
import shutil
from collections import Counter
from datetime import date as date_cls, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
SKILL = BASE.parent
DEFAULT_TEMPLATE = SKILL / "assets" / "report_template.html"
PHONE_RE = re.compile(r"^\+?\d[\d\s-]{6,}$")
WEEKDAYS = ["星期一", "星期二", "星期三", "星期四", "星期五", "星期六", "星期日"]


def parse_dt(value: str | None) -> datetime | None:
    if not value:
        return None
    try:
        return datetime.fromisoformat(value)
    except ValueError:
        return None


def hm(value: datetime | None) -> str:
    return value.strftime("%H:%M") if value else "-"


def friendly_group_name(name: str) -> str:
    text = re.sub(r"^[0-9\ufe0f\u20e3\u200b\s]+", "", name)
    text = text.replace("🍠", "").strip()
    return text or name


def mask_contact(raw: str, mask: bool) -> str:
    value = str(raw or "").strip()
    if not mask:
        return value
    if PHONE_RE.match(value):
        digits = re.sub(r"\D", "", value)
        if len(digits) >= 7:
            return f"{digits[:3]}****{digits[-4:]}"
    if len(value) <= 3:
        return value[:1] + "***"
    if len(value) <= 6:
        return value[:2] + "***"
    return value[:4] + "***" + value[-2:]


def contact_type(raw: str) -> str:
    value = str(raw or "").strip()
    if PHONE_RE.match(value) and len(re.sub(r"\D", "", value)) >= 10:
        return "phone"
    return "wechat"


def is_image(record: dict[str, Any]) -> bool:
    class_name = str(record.get("class") or "")
    text = str(record.get("content") or "").strip()
    return "Image" in class_name or text in {"图片", "[图片]"}


def resolve_date(raw: str, now: datetime | None = None) -> date_cls:
    current = now or datetime.now()
    text = (raw or "yesterday").strip().lower()
    if text in {"yesterday", "昨天"}:
        return current.date() - timedelta(days=1)
    if text in {"today", "今天"}:
        return current.date()
    return datetime.strptime(raw, "%Y-%m-%d").date()


def load_messages(source: Path) -> dict[str, list[dict[str, Any]]]:
    messages: dict[str, list[dict[str, Any]]] = {}
    for path in source.glob("messages_*.jsonl"):
        with path.open(encoding="utf-8") as handle:
            for line in handle:
                if not line.strip():
                    continue
                record = json.loads(line)
                messages.setdefault(record["group"], []).append(record)
    for records in messages.values():
        records.sort(key=lambda item: item["seq"])
    return messages


def new_lead(group_name: str, record: dict[str, Any]) -> dict[str, Any]:
    value = parse_dt(record.get("resolved"))
    return {
        "group": group_name,
        "start": value,
        "end": value,
        "last_dt": value,
        "publisher_set": [],
        "contacts": [],
        "images": [],
        "image_count": 0,
        "claimers": [],
        "evidence": [],
        "followups": [],
        "seq_start": record["seq"],
        "seq_end": record["seq"],
        "has_claim": False,
    }


def add_publisher_message(lead: dict[str, Any], record: dict[str, Any], mask: bool) -> None:
    value = parse_dt(record.get("resolved"))
    if value:
        lead["last_dt"] = value
        lead["end"] = value
    lead["seq_end"] = record["seq"]
    sender = str(record.get("sender") or "")
    if sender and sender not in lead["publisher_set"]:
        lead["publisher_set"].append(sender)
    if is_image(record):
        image_path = str(record.get("image_path") or "")
        lead["image_count"] += 1
        if image_path:
            lead["images"].append(image_path)
        lead["evidence"].append(
            {
                "seq": record["seq"],
                "time": value.isoformat() if value else None,
                "time_label": hm(value),
                "sender": sender,
                "text": "[图片]" if not image_path else f"[图片] {image_path}",
                "kind": "image",
            }
        )
        return
    raw = str(record.get("content") or "").strip()
    display = mask_contact(raw, mask)
    lead["contacts"].append({"type": contact_type(raw), "display": display})
    lead["evidence"].append(
        {
            "seq": record["seq"],
            "time": value.isoformat() if value else None,
            "time_label": hm(value),
            "sender": sender,
            "text": display,
            "kind": "contact",
        }
    )


def materialize_lead_images(
    leads: list[dict[str, Any]], source: Path, image_out: Path | None
) -> None:
    """把 probe 截下的图片复制进报告目录，并补上可读的绝对路径。"""
    if image_out is None:
        return
    for lead in leads:
        original = list(lead.get("images") or [])
        saved_paths: list[str] = []
        saved_abs: list[str] = []
        image_rows = [row for row in lead.get("evidence", []) if row.get("kind") == "image"]
        for index, relative in enumerate(original):
            src = source / relative
            report_path = ""
            absolute = ""
            if src.exists():
                dest_dir = image_out / "images"
                dest_dir.mkdir(parents=True, exist_ok=True)
                dest = dest_dir / src.name
                if not dest.exists():
                    shutil.copy2(src, dest)
                report_path = f"images/{src.name}"
                absolute = str(dest.resolve())
                saved_paths.append(report_path)
                saved_abs.append(absolute)
            if index < len(image_rows):
                row = image_rows[index]
                row["text"] = "[图片]"
                row["source_image"] = relative
                if report_path:
                    row["image"] = report_path
                    row["image_abs"] = absolute
        lead["images"] = saved_paths
        lead["image_abs"] = saved_abs
        lead["image_saved"] = len(saved_paths)


def build_group(
    group_name: str,
    records: list[dict[str, Any]],
    claims: dict[int, dict[str, Any]],
    publishers: list[str],
    batch_minutes: int,
    mask: bool,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    leads: list[dict[str, Any]] = []
    orphans: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None

    def close() -> None:
        nonlocal current
        if current and (current["contacts"] or current["image_count"]):
            leads.append(current)
        current = None

    for record in records:
        if record.get("is_time") or record.get("is_system"):
            continue
        seq = record["seq"]
        sender = str(record.get("sender") or "")
        value = parse_dt(record.get("resolved"))
        class_name = str(record.get("class") or "")
        claim = claims.get(seq)

        if claim:
            if current is None:
                previous = None
                join_record = None
                for item in records:
                    if item["seq"] >= seq:
                        break
                    previous = item
                    text = str(item.get("content") or "")
                    if item.get("is_system") and ("加入群聊" in text or "邀请" in text):
                        join_record = item
                evidence = [
                    {
                        "seq": seq,
                        "time": value.isoformat() if value else None,
                        "time_label": hm(value),
                        "sender": sender,
                        "text": "1",
                        "kind": "claim",
                    }
                ]
                context_record = join_record or (
                    previous if previous and previous.get("is_system") else None
                )
                if context_record:
                    text = str(context_record.get("content") or "")
                    evidence.insert(
                        0,
                        {
                            "seq": context_record["seq"],
                            "time": context_record.get("resolved"),
                            "time_label": hm(parse_dt(context_record.get("resolved"))),
                            "sender": "系统",
                            "text": text[:120] + ("…" if len(text) > 120 else ""),
                            "kind": "system",
                        },
                    )
                orphans.append(
                    {
                        "id": f"X{len(orphans) + 1:02d}",
                        "kind": "orphan",
                        "group": group_name,
                        "time": value.isoformat() if value else None,
                        "time_label": hm(value),
                        "sender": sender,
                        "status": "orphan",
                        "note": "前面没有可见线索；可能是本机登录前的消息，或发布人未配置。",
                        "evidence": evidence,
                    }
                )
            else:
                order = len(current["claimers"]) + 1
                current["claimers"].append(
                    {
                        "name": sender,
                        "time": value.isoformat() if value else None,
                        "time_label": hm(value),
                        "order": order,
                    }
                )
                current["evidence"].append(
                    {
                        "seq": seq,
                        "time": value.isoformat() if value else None,
                        "time_label": hm(value),
                        "sender": sender,
                        "text": "1",
                        "kind": "claim",
                    }
                )
                current["has_claim"] = True
            continue

        if sender in publishers and "Quote" not in class_name:
            if current is None:
                current = new_lead(group_name, record)
                add_publisher_message(current, record, mask)
                continue
            gap = (
                (value - current["last_dt"]).total_seconds() / 60
                if value and current["last_dt"]
                else 999
            )
            if current["has_claim"] or gap > batch_minutes:
                close()
                current = new_lead(group_name, record)
                add_publisher_message(current, record, mask)
            else:
                add_publisher_message(current, record, mask)
            continue

        if "Quote" in class_name:
            continue
        if current and value and current["last_dt"]:
            gap = (value - current["last_dt"]).total_seconds() / 60
            if gap <= batch_minutes:
                text = str(record.get("content") or "")
                if any(word in text for word in ("已加", "加了", "通过", "拉群", "带看", "联系")):
                    current["followups"].append(
                        {"sender": sender, "time_label": hm(value), "text": text}
                    )
                current["evidence"].append(
                    {
                        "seq": seq,
                        "time": value.isoformat() if value else None,
                        "time_label": hm(value),
                        "sender": sender,
                        "text": text,
                        "kind": "other",
                    }
                )

    close()
    return leads, orphans


def finalize_lead(lead: dict[str, Any]) -> dict[str, Any]:
    lead["time"] = lead["start"].isoformat() if lead["start"] else None
    lead["time_label"] = hm(lead["start"])
    lead["publisher"] = " / ".join(lead["publisher_set"])
    lead["contact_count"] = len(lead["contacts"])
    lead["status"] = "claimed" if lead["claimers"] else "unclaimed"
    lead["claim_count"] = len(lead["claimers"])
    lead["first_claimer"] = lead["claimers"][0]["name"] if lead["claimers"] else ""
    lead["claim_time_label"] = lead["claimers"][0]["time_label"] if lead["claimers"] else ""
    lead.pop("start", None)
    lead.pop("end", None)
    lead.pop("last_dt", None)
    lead.pop("has_claim", None)
    return lead


def load_config(path: Path) -> dict[str, Any]:
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    groups = data.get("groups") or []
    if not groups:
        raise ValueError("groups.json 里没有 groups")
    for index, group in enumerate(groups, start=1):
        group.setdefault("key", f"g{index}")
        group.setdefault("short", f"群{index}")
        group.setdefault("display_name", friendly_group_name(group["name"]))
        group.setdefault("publishers", [])
    data["batch_minutes"] = int(data.get("batch_minutes") or 5)
    data["mask"] = bool(data.get("mask", False))
    return data


def load_data(
    source: Path,
    config: dict[str, Any],
    date_value: date_cls,
    mask: bool,
    image_out: Path | None = None,
) -> dict[str, Any]:
    report = json.loads((source / "report.json").read_text(encoding="utf-8"))
    messages = load_messages(source)
    sample_tz = None
    for records in messages.values():
        for record in records:
            value = parse_dt(record.get("resolved"))
            if value is not None:
                sample_tz = value.tzinfo
                break
        if sample_tz is not None:
            break
    start = datetime.combine(date_value, datetime.min.time(), tzinfo=sample_tz)
    end = start + timedelta(days=1)
    for group_name, records in messages.items():
        messages[group_name] = [
            record
            for record in records
            if (value := parse_dt(record.get("resolved"))) and start <= value < end
        ]
    report_by_name = {item["group"]: item for item in report.get("groups", [])}
    config_by_name = {item["name"]: item for item in config["groups"]}
    batch_minutes = config["batch_minutes"]

    groups: list[dict[str, Any]] = []
    all_leads: list[dict[str, Any]] = []
    all_orphans: list[dict[str, Any]] = []
    issue_groups: list[dict[str, Any]] = []

    ordered_names: list[str] = [item["name"] for item in config["groups"]]
    for name in report_by_name:
        if name not in ordered_names:
            ordered_names.append(name)

    for index, name in enumerate(ordered_names, start=1):
        cfg = config_by_name.get(name, {})
        group_data = report_by_name.get(name)
        short = str(cfg.get("short") or f"群{index}")
        display_name = str(cfg.get("display_name") or friendly_group_name(name))
        group_label = f"{short} · {display_name}"
        publishers = list(cfg.get("publishers") or [])
        error = ""

        if group_data is None:
            error = "未采集：群名没有出现在本次 probe 结果里"
            group_data = {
                "group": name,
                "chat_info": {},
                "claims": [],
                "gaps": [],
                "earliest": None,
                "latest": None,
                "merged_count": 0,
                "in_range_count": 0,
                "reached_since": False,
                "bottom_reached": False,
                "error": error,
            }
        else:
            error = str(group_data.get("error") or "")
            if "连续多轮没有新增" in error:
                error = ""
        if not publishers:
            error = error or "未配置发布人，无法切分线索"

        claims = {claim["seq"]: claim for claim in group_data.get("claims", [])}
        leads, orphans = build_group(
            name,
            messages.get(name, []),
            claims,
            publishers,
            batch_minutes,
            mask,
        )
        leads = [finalize_lead(lead) for lead in leads]
        materialize_lead_images(leads, source, image_out)
        for lead in leads:
            lead["kind"] = "lead"
            lead["group_key"] = cfg.get("key") or f"g{index}"
            lead["group_short"] = short
            lead["group_display"] = display_name
            lead["group_label"] = group_label
            if lead["image_count"]:
                saved = int(lead.get("image_saved") or 0)
                if saved >= lead["image_count"]:
                    lead["note"] = f"含 {lead['image_count']} 张图片，点“查看原文”可看。"
                elif saved:
                    lead["note"] = f"含 {lead['image_count']} 张图片，已保存 {saved} 张。"
                else:
                    lead["note"] = f"含 {lead['image_count']} 张图片；图片未落盘，只能看到占位。"
            elif lead["status"] == "unclaimed":
                lead["note"] = ""
            else:
                lead["note"] = ""

        for orphan in orphans:
            orphan["group_key"] = cfg.get("key") or f"g{index}"
            orphan["group_short"] = short
            orphan["group_display"] = display_name
            orphan["group_label"] = group_label

        claimers = Counter(claimer["name"] for lead in leads for claimer in lead["claimers"])
        claimers.update(orphan["sender"] for orphan in orphans)
        unclaimed_contacts = sum(
            lead["contact_count"] for lead in leads if lead["status"] == "unclaimed"
        )
        earliest = parse_dt(group_data.get("earliest"))
        latest = parse_dt(group_data.get("latest"))
        gaps = len(group_data.get("gaps") or [])
        reached_since = bool(group_data.get("reached_since"))
        image_count = sum(lead["image_count"] for lead in leads)
        is_issue = bool(error or gaps or not reached_since)
        if error:
            status = "review"
            status_label = "需核对"
        elif gaps:
            status = "review"
            status_label = "需核对"
        elif not reached_since:
            status = "partial"
            status_label = "部分"
        else:
            status = "complete"
            status_label = "完整"

        status_bits = []
        if error:
            status_bits.append(error)
        if not reached_since:
            status_bits.append(f"只覆盖 {hm(earliest)} 以后")
        if gaps:
            status_bits.append(f"{gaps} 处低置信")
        status_note = "；".join(status_bits) or "合并无缺口"

        group = {
            "key": cfg.get("key") or f"g{index}",
            "name": name,
            "display_name": display_name,
            "label": group_label,
            "short": short,
            "members": int((group_data.get("chat_info") or {}).get("group_member_count") or 0),
            "leads": len(leads),
            "claimed": sum(1 for lead in leads if lead["status"] == "claimed"),
            "unclaimed": sum(1 for lead in leads if lead["status"] == "unclaimed"),
            "claimers": len(claimers),
            "claims": sum(len(lead["claimers"]) for lead in leads) + len(orphans),
            "orphans": len(orphans),
            "unclaimed_contacts": unclaimed_contacts,
            "coverage_from": hm(earliest),
            "coverage_to": hm(latest),
            "earliest": group_data.get("earliest"),
            "latest": group_data.get("latest"),
            "merged_count": int(group_data.get("merged_count") or 0),
            "in_range_count": int(group_data.get("in_range_count") or 0),
            "gaps": gaps,
            "reached_since": reached_since,
            "image_count": image_count,
            "image_saved": sum(int(lead.get("image_saved") or 0) for lead in leads),
            "status": status,
            "status_label": status_label,
            "status_note": status_note,
            "bar": round(
                sum(1 for lead in leads if lead["status"] == "claimed") / len(leads) * 100
            )
            if leads
            else 0,
            "error": error,
        }
        groups.append(group)
        if is_issue:
            issue_groups.append(group)
        all_leads.extend(leads)
        all_orphans.extend(orphans)

    all_leads.sort(key=lambda lead: lead["time"] or "")
    all_orphans.sort(key=lambda orphan: orphan["time"] or "")
    for index, lead in enumerate(all_leads, start=1):
        lead["id"] = f"L{index:02d}"
    for index, orphan in enumerate(all_orphans, start=1):
        orphan["id"] = f"X{index:02d}"

    lead_count = len(all_leads)
    claimed = sum(1 for lead in all_leads if lead["status"] == "claimed")
    unclaimed = sum(1 for lead in all_leads if lead["status"] == "unclaimed")
    claims = sum(len(lead["claimers"]) for lead in all_leads) + len(all_orphans)
    unclaimed_contacts = sum(
        lead["contact_count"] for lead in all_leads if lead["status"] == "unclaimed"
    )
    top_claimers = Counter(
        claimer["name"] for lead in all_leads for claimer in lead["claimers"]
    )
    top_claimers.update(orphan["sender"] for orphan in all_orphans)
    issue_names = "、".join(group["short"] for group in issue_groups) or "无"

    target_since = parse_dt(report.get("since"))
    target_until = parse_dt(report.get("until"))
    now = datetime.now()
    date_label = f"{date_value.year} 年 {date_value.month} 月 {date_value.day} 日 · {WEEKDAYS[date_value.weekday()]}"
    hero_prefix = (
        "昨天，"
        if date_value == now.date() - timedelta(days=1)
        else f"{date_value.month} 月 {date_value.day} 日，"
    )
    visible_parts = [
        f"{group['short']} {group['coverage_from']}–{group['coverage_to']}"
        for group in groups
        if group["coverage_from"] != "-" and group["coverage_to"] != "-"
    ]
    scope_visible = (
        "本机实际可见：" + "、".join(visible_parts)
        if any(not group["reached_since"] for group in groups)
        else "本机可见范围覆盖目标起点"
    )
    if not visible_parts:
        scope_visible = "本机没有读到任何群消息"

    unclaimed_leads = [lead for lead in all_leads if lead["status"] == "unclaimed"]
    attention: list[dict[str, Any]] = []
    if unclaimed_leads:
        unclaimed_groups = [
            group for group in groups if group["unclaimed"] > 0
        ]
        times = "、".join(lead["time_label"] for lead in unclaimed_leads[:3])
        if len(unclaimed_groups) == 1:
            title = (
                f"{unclaimed_groups[0]['short']} · {unclaimed_groups[0]['unclaimed']} 批共 "
                f"{unclaimed_groups[0]['unclaimed_contacts']} 个联系方式没人扣 1"
            )
        else:
            title = (
                f"{len(unclaimed_groups)} 个群 · {len(unclaimed_leads)} 批共 "
                f"{unclaimed_contacts} 个联系方式没人扣 1"
            )
        attention.append(
            {
                "kind": "unclaimed",
                "title": title,
                "detail": f"{times}；可能私聊接了，群里看不到。",
                "action": "看明细",
                "target": "unclaimed",
            }
        )
    if all_orphans:
        first = all_orphans[0]
        title = (
            f"{first['group_short']} · {first['time_label']} 有一条“1”找不到对应线索"
            if len(all_orphans) == 1
            else f"{len(all_orphans)} 条“1”找不到对应线索"
        )
        attention.append(
            {
                "kind": "orphan",
                "title": title,
                "detail": "本机可见起点之前没有线索，或发布人未配置；先人工确认。",
                "action": "看证据",
                "target": first["id"],
            }
        )
    if issue_groups:
        attention.append(
            {
                "kind": "review",
                "title": f"{len(issue_groups)} 个群数据需核对：{issue_names}",
                "detail": "覆盖不完整、合并低置信或群名/发布人未配置；详见数据说明。",
                "action": "看数据说明",
                "target": "notes",
            }
        )
    if not attention:
        attention.append(
            {
                "kind": "good",
                "title": "昨天没有需要处理的异常",
                "detail": "所有群都覆盖目标起点，扣 1 都能对上；仍建议抽查原文。",
                "action": "看明细",
                "target": "detail",
            }
        )

    insights: list[str] = []
    if top_claimers:
        top = top_claimers.most_common(3)
        insights.append(
            "扣得最多的是 "
            + "、".join(f"{name} {count} 次" for name, count in top)
            + "。"
        )
    else:
        insights.append("昨天还没有人扣 1。")
    if unclaimed_leads:
        insights.append(
            f"{len(unclaimed_leads)} 批联系方式没人扣 1，共 {unclaimed_contacts} 个联系方式；可能私聊接了。"
        )
    else:
        insights.append("所有线索在群里都有人扣 1。")
    followup_count = sum(len(lead["followups"]) for lead in all_leads)
    if followup_count:
        insights.append(f"群内可见 {followup_count} 条后续动作；仍以群内消息为准。")
    else:
        insights.append("群里没有看到“已加/拉群/带看”的后续动作；只能确认谁扣了 1，不能确认谁接上了。")
    if issue_groups:
        insights.append(f"{len(issue_groups)} 个群数据需核对：{issue_names}。")

    notes = [
        {
            "label": "统计范围",
            "value": f"{target_since:%Y-%m-%d %H:%M}–{target_until:%Y-%m-%d %H:%M}"
            if target_since and target_until
            else "目标区间未知",
        },
        {
            "label": "本机实际可见",
            "value": "；".join(
                f"{group['short']} {group['coverage_from']}–{group['coverage_to']}"
                if group["coverage_from"] != "-"
                else f"{group['short']} 未读到消息"
                for group in groups
            ),
        },
        {
            "label": "为什么不是完整一天",
            "value": (
                "采集电脑在 "
                + "、".join(
                    f"{group['short']} {group['coverage_from']}"
                    for group in groups
                    if not group["reached_since"] and group["coverage_from"] != "-"
                )
                + " 才可见；登录前的群历史读不到。"
                if any(not group["reached_since"] for group in groups)
                else "本次可见范围覆盖目标起点；但本机只读本地已有记录，仍以实际可见为准。"
            ),
        },
        {
            "label": "只统计群内",
            "value": "只统计群里可见的扣 1；私聊领取、其他群里的动作看不到。",
        },
        {
            "label": "数据质量",
            "value": "；".join(
                f"{group['short']} 合并 {group['merged_count']} 条"
                + (f"，{group['gaps']} 处低置信" if group["gaps"] else "，无缺口")
                + (
                    f"，{group['image_count']} 张图片，已落盘 {group.get('image_saved', 0)} 张"
                    if group["image_count"]
                    else ""
                )
                + (f"，{group['error']}" if group["error"] else "")
                for group in groups
            ),
        },
        {
            "label": "线索口径",
            "value": f"同一发布人 {batch_minutes} 分钟内连续发、中间没人扣 1 的联系方式，算一个线索批次；有人扣 1 就切分。",
        },
        {
            "label": "发布人配置",
            "value": "发布人来自 config/groups.json；未配置的群会标异常，不硬猜。",
        },
        {
            "label": "群名显示",
            "value": "采集用原始群名；页面用 display_name 简化。",
        },
        {
            "label": "数据来源",
            "value": f"{source}；" + ("联系方式已脱敏" if mask else "联系方式未打码，文件只留本机"),
        },
    ]

    meta = {
        "title": f"{date_value.month} 月 {date_value.day} 日 · 群日报",
        "date": date_value.isoformat(),
        "date_label": date_label,
        "generated_at": now.strftime("%Y-%m-%d %H:%M"),
        "hero_prefix": hero_prefix,
        "brand_sub": f"{len(groups)} 个群 · 只读采集",
        "badge": "内部日报",
        "eyebrow": f"{date_label} · 群日报",
        "scope_target": f"{date_value.month}-{date_value.day:02d} 00:00–24:00",
        "scope_visible": scope_visible,
        "groups_foot": f"{len(groups)} 个群 · 已按发布批次合并",
        "claims_foot": f"共 {claims} 条扣 1 记录",
        "unclaimed_foot": (
            f"{unclaimed} 批，共 {unclaimed_contacts} 个联系方式"
            if unclaimed
            else "没有没人扣 1 的批次"
        ),
        "review_foot": (
            f"{len(all_orphans)} 条无对应线索 · {len(issue_groups)} 个群数据需核对"
        ),
        "groups_title": f"{len(groups)} 个群",
        "insights_title": f"豆包看到的 {len(insights)} 件事",
        "disclaimer": f"本次共 {len(groups)} 个群；只统计群内可见的扣 1；私聊领取、其他群动作看不到。",
        "footer_source": f"生成于 {now:%Y-%m-%d %H:%M} · 数据源 {source.name}",
        "footer_note": "只读采集 · 不发送消息 · " + ("联系方式已脱敏" if mask else "数据只留本机"),
        "export_name": f"群日报_{date_value.isoformat()}.xlsx",
        "lede": (
            f"共 {len(groups)} 个群、{lead_count} 条线索；"
            f"{claims} 条扣 1 记录，{len(all_orphans)} 条没有对应线索。"
            "先看下面几件事，明细点开再看。"
        ),
    }

    return {
        "meta": meta,
        "stats": {
            "groups": len(groups),
            "leads": lead_count,
            "claimed": claimed,
            "unclaimed": unclaimed,
            "claims": claims,
            "orphans": len(all_orphans),
            "review": len(all_orphans) + len(issue_groups),
            "unclaimed_contacts": unclaimed_contacts,
        },
        "groups": groups,
        "leads": all_leads,
        "orphans": all_orphans,
        "attention": attention,
        "insights": insights,
        "notes": notes,
    }


def style_sheet(sheet, headers: list[str], widths: list[int], rows: list[list[Any]]) -> None:
    header_fill = PatternFill("solid", fgColor="0E7C66")
    header_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
    body_font = Font(name="微软雅黑", size=10, color="1F2826")
    thin = Side(style="thin", color="E5E0D6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    band_fill = PatternFill("solid", fgColor="FBF9F4")
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    sheet.row_dimensions[1].height = 28
    for row_index, row in enumerate(rows, start=2):
        sheet.append(row)
        for column_index, cell in enumerate(sheet[row_index], start=1):
            cell.font = body_font
            cell.border = border
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=column_index in {2, 6, 8, 12, 13},
            )
            if row_index % 2 == 0:
                cell.fill = band_fill
        sheet.row_dimensions[row_index].height = 24
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(sheet.max_row, 1)}"
    sheet.sheet_view.showGridLines = False


def build_xlsx(data: dict[str, Any], path: Path) -> None:
    meta = data["meta"]
    workbook = Workbook()
    workbook.properties.title = meta["title"]
    workbook.properties.creator = "wechat-group-daily"

    detail_sheet = workbook.active
    detail_sheet.title = "线索明细"
    detail_rows = []
    for lead in data["leads"]:
        contacts = "、".join(contact["display"] for contact in lead["contacts"])
        if len(lead["contacts"]) > 8:
            contacts = "、".join(contact["display"] for contact in lead["contacts"][:8])
            contacts += f" 等 {len(lead['contacts'])} 个"
        claimers = "；".join(
            f"{claimer['name']}（第 {claimer['order']} 个，{claimer['time_label']}）"
            for claimer in lead["claimers"]
        )
        detail_rows.append(
            [
                meta["date"],
                lead["group_label"],
                lead["id"],
                lead["publisher"],
                lead["time_label"],
                contacts or ("[图片]" if lead["image_count"] else ""),
                lead["image_count"],
                claimers or "—",
                lead["claim_time_label"] or "—",
                "第 1 个" if lead["claimers"] else "—",
                "有人扣 1" if lead["status"] == "claimed" else "没人扣 1",
                "群内暂无后续动作" if not lead["followups"] else "；".join(item["text"] for item in lead["followups"]),
                lead.get("note") or "",
            ]
        )
    style_sheet(
        detail_sheet,
        ["日期", "群", "线索编号", "发布人", "发布时间", "联系方式", "图片数", "扣1人", "扣1时间", "顺序", "状态", "后续动作", "备注"],
        [12, 28, 9, 22, 10, 34, 7, 26, 10, 8, 11, 20, 34],
        detail_rows,
    )
    claimed_fill = PatternFill("solid", fgColor="E6F3EF")
    unclaimed_fill = PatternFill("solid", fgColor="FDF0DC")
    for row_index in range(2, detail_sheet.max_row + 1):
        cell = detail_sheet.cell(row=row_index, column=11)
        cell.fill = claimed_fill if cell.value == "有人扣 1" else unclaimed_fill
        cell.font = Font(name="微软雅黑", bold=True, size=10, color="0A5C4C" if cell.value == "有人扣 1" else "B45309")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    detail_sheet.sheet_properties.tabColor = "0E7C66"

    claim_sheet = workbook.create_sheet("领取记录")
    claim_rows = []
    for lead in data["leads"]:
        for claimer in lead["claimers"]:
            claim_rows.append(
                [
                    meta["date"],
                    lead["group_label"],
                    lead["id"],
                    claimer["name"],
                    claimer["time_label"],
                    claimer["order"],
                    "对应线索",
                    "正常",
                ]
            )
    for orphan in data["orphans"]:
        claim_rows.append(
            [
                meta["date"],
                orphan["group_label"],
                orphan["id"],
                orphan["sender"],
                orphan["time_label"],
                1,
                "没有对应线索",
                "需核对",
            ]
        )
    claim_rows.sort(key=lambda row: (row[4], row[1], row[3]))
    style_sheet(
        claim_sheet,
        ["日期", "群", "线索编号", "扣1人", "扣1时间", "顺序", "对应线索", "状态"],
        [12, 28, 10, 22, 10, 8, 18, 11],
        claim_rows,
    )
    review_fill = PatternFill("solid", fgColor="FDE8E6")
    for row_index in range(2, claim_sheet.max_row + 1):
        cell = claim_sheet.cell(row=row_index, column=8)
        if cell.value == "需核对":
            cell.fill = review_fill
            cell.font = Font(name="微软雅黑", bold=True, size=10, color="B42318")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    claim_sheet.sheet_properties.tabColor = "B45309"

    review_sheet = workbook.create_sheet("异常")
    review_rows = []
    for orphan in data["orphans"]:
        review_rows.append(
            [
                orphan["group_label"],
                "无对应线索",
                f"{orphan['time_label']} 有一条领取动作，前面没有可见线索。",
                orphan["id"],
            ]
        )
    for group in data["groups"]:
        if group["error"]:
            review_rows.append([group["label"], "采集异常", group["error"], ""])
        if not group["reached_since"]:
            review_rows.append(
                [group["label"], "覆盖不完整", f"本机可见 {group['coverage_from']}–{group['coverage_to']}，目标区间更早。", "数据说明"]
            )
        if group["gaps"]:
            review_rows.append(
                [group["label"], "低置信合并", f"合并 {group['merged_count']} 条，{group['gaps']} 处低置信。", "gaps.json"]
            )
        if group["image_count"]:
            saved = int(group.get("image_saved") or 0)
            if saved == 0:
                review_rows.append(
                    [group["label"], "图片未落盘", f"{group['image_count']} 张图片未保存原图；页面只显示占位。", ""]
                )
            elif saved < group["image_count"]:
                review_rows.append(
                    [
                        group["label"],
                        "图片部分未落盘",
                        f"{group['image_count']} 张图片，已保存 {saved} 张，其余只显示占位。",
                        "",
                    ]
                )
        low = group.get("low_confidence_count") or 0
        if low:
            review_rows.append(
                [group["label"], "低置信线索", f"{low} 条线索 AI 置信度低，建议点开原文核对。", ""]
            )
    if not review_rows:
        review_rows.append(["全部群", "无异常", "本次没有发现采集或扣 1 异常。", ""])
    style_sheet(
        review_sheet,
        ["群", "类型", "说明", "证据"],
        [30, 16, 70, 18],
        review_rows,
    )
    review_sheet.sheet_properties.tabColor = "B42318"

    notes_sheet = workbook.create_sheet("数据说明")
    notes_sheet.append(["项目", "说明"])
    header_fill = PatternFill("solid", fgColor="0E7C66")
    header_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
    for cell in notes_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for note in data["notes"]:
        notes_sheet.append([note["label"], note["value"]])
    for row_index in range(2, notes_sheet.max_row + 1):
        for column_index in range(1, 3):
            cell = notes_sheet.cell(row=row_index, column=column_index)
            cell.font = Font(name="微软雅黑", size=10, color="1F2826")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        notes_sheet.row_dimensions[row_index].height = 34
    notes_sheet.column_dimensions["A"].width = 22
    notes_sheet.column_dimensions["B"].width = 90
    notes_sheet.freeze_panes = "A2"
    notes_sheet.sheet_view.showGridLines = False
    notes_sheet.sheet_properties.tabColor = "6B746F"

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def build_html(data: dict[str, Any], template_path: Path, path: Path) -> None:
    template = template_path.read_text(encoding="utf-8")
    payload = json.dumps(data, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    meta = data["meta"]
    replacements = {
        "__DATA__": payload,
        "__TITLE__": html.escape(meta["title"]),
        "__BRAND_SUB__": html.escape(meta["brand_sub"]),
        "__BADGE__": html.escape(meta["badge"]),
        "__EYEBROW__": html.escape(meta["eyebrow"]),
        "__HERO_PREFIX__": html.escape(meta["hero_prefix"]),
        "__SCOPE_TARGET__": html.escape(meta["scope_target"]),
        "__SCOPE_VISIBLE__": html.escape(meta["scope_visible"]),
        "__GROUPS_FOOT__": html.escape(meta["groups_foot"]),
        "__CLAIMS_FOOT__": html.escape(meta["claims_foot"]),
        "__UNCLAIMED_FOOT__": html.escape(meta["unclaimed_foot"]),
        "__REVIEW_FOOT__": html.escape(meta["review_foot"]),
        "__GROUPS_TITLE__": html.escape(meta["groups_title"]),
        "__DISCLAIMER__": html.escape(meta["disclaimer"]),
        "__INSIGHTS_TITLE__": html.escape(meta["insights_title"]),
        "__FOOTER_SOURCE__": html.escape(meta["footer_source"]),
        "__FOOTER_NOTE__": html.escape(meta["footer_note"]),
        "__GENERATED_AT__": html.escape(meta["generated_at"]),
        "__EXPORT_NAME__": html.escape(meta["export_name"]),
    }
    for key, value in replacements.items():
        template = template.replace(key, value)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(template, encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description="从 probe 输出生成群日报")
    parser.add_argument("--source", required=True, help="probe 输出目录")
    parser.add_argument("--config", required=True, help="groups.json")
    parser.add_argument("--out", required=True, help="报告输出目录")
    parser.add_argument("--date", default="yesterday", help="yesterday / today / 2026-10-28")
    parser.add_argument("--mask", action=argparse.BooleanOptionalAction, default=None, help="是否打码；默认读 config.mask，通常为 false")
    parser.add_argument("--template", default=str(DEFAULT_TEMPLATE), help="HTML 模板")
    args = parser.parse_args()

    source = Path(args.source).resolve()
    config_path = Path(args.config).resolve()
    out = Path(args.out).resolve()
    config = load_config(config_path)
    date_value = resolve_date(args.date)
    mask = config["mask"] if args.mask is None else bool(args.mask)
    out.mkdir(parents=True, exist_ok=True)
    data = load_data(source, config, date_value, mask, image_out=out)
    (out / "facts.json").write_text(
        json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    build_xlsx(data, out / data["meta"]["export_name"])
    build_html(data, Path(args.template), out / "report.html")
    print(
        json.dumps(
            {
                "date": data["meta"]["date"],
                "stats": data["stats"],
                "report": str(out / "report.html"),
                "xlsx": str(out / data["meta"]["export_name"]),
                "facts": str(out / "facts.json"),
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
